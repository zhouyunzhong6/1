"""生成数据底表（xlsx + csv）与附录 Markdown 片段"""
import csv
import json
import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

from db_market import COMPANIES, FIN_EVENTS, FIN_STATS, SHIPMENTS
from db_sources import SOURCES
from db_dd import DD
from db_theses import THESES
from db_matrices import TECH_MATRIX, SCENARIOS, OPPORTUNITIES, SCORECARD, VALUE_CHAIN

HERE = os.path.dirname(os.path.abspath(__file__))
GEN = os.path.join(HERE, "..", "generated")
DATA = os.path.join(HERE, "..", "data")
os.makedirs(DATA, exist_ok=True)

with open(os.path.join(GEN, "model_results.json"), encoding="utf-8") as f:
    MODEL = json.load(f)

SHEETS = []


def add(name, header, rows, widths=None, note=None):
    SHEETS.append((name, header, rows, widths, note))


def md_table(header, rows):
    esc = lambda s: str(s).replace("|", "／").replace("\n", " ")
    out = ["| " + " | ".join(header) + " |", "|" + "---|" * len(header)]
    for r in rows:
        out.append("| " + " | ".join(esc(c) for c in r) + " |")
    return "\n".join(out)


def save_md(name, text):
    with open(os.path.join(GEN, name + ".md"), "w", encoding="utf-8") as f:
        f.write(text.strip() + "\n")


# ---------------- 组装各表 ----------------
H_COMP = ["名称", "地区", "产业链环节", "核心产品/路线", "资本状态", "最新估值/市值（时间）", "收入/出货（最新）", "成熟度", "证据等级", "关键风险/备注"]
add("01_企业图谱", H_COMP, COMPANIES, [22, 10, 16, 30, 22, 30, 44, 14, 8, 36],
    "成熟度：L1 演示/L2 受控持续运行/L3 客户试点/L4 可重复付费部署/L5 规模化成熟产品；证据：A 一手/B 可靠二手交叉佐证/C 单方或单一来源/D 推断")
H_FIN = ["日期", "公司", "地区", "事件/轮次", "金额", "估值", "状态", "证据等级"]
add("02_融资与资本事件", H_FIN, FIN_EVENTS, [12, 18, 8, 28, 26, 26, 16, 8], "状态严格区分：已完成/已宣布/传闻/审核中；同一轮次仅计一次")
H_STAT = ["范围", "统计机构", "期间", "笔数", "金额", "口径说明", "证据等级"]
add("03_融资统计口径对比", H_STAT, FIN_STATS, [16, 22, 18, 20, 18, 44, 8], "不同机构口径差异大，不可相加或直接比较")
H_SHIP = ["来源", "期间", "范围", "数值", "口径", "证据等级"]
add("04_出货与产量口径", H_SHIP, SHIPMENTS, [26, 16, 18, 50, 40, 8], "产量≠出货≠销量≠部署≠有效运行")
H_TECH = ["类别", "技术/部件", "成熟度", "关键指标现状（2026-10）", "主要瓶颈", "进步速度", "接近物理/工程约束", "降本空间", "投资含义"]
add("05_技术成熟度矩阵", H_TECH, TECH_MATRIX, [8, 28, 16, 40, 30, 8, 14, 10, 34])
H_SCEN = ["场景", "典型任务", "现有方案", "人形必要性", "商业化门槛", "2026 阶段", "回本周期", "决策方", "优先级"]
add("06_场景商业化比较", H_SCEN, SCENARIOS, [22, 26, 24, 18, 28, 18, 26, 16, 14])
H_OPP = ["方向", "环节", "产业阶段", "确定性", "潜在回报", "估值热度", "被替代风险", "优先级", "建议阶段与估值", "进入条件", "放弃条件"]
add("07_投资机会优先级", H_OPP, OPPORTUNITIES, [30, 10, 16, 8, 8, 10, 16, 10, 30, 24, 24])
H_SC = ["维度", "精密部件", "场景闭环机器人", "整机平台", "基础模型", "数据/工具链"]
add("08_评分卡权重", H_SC, SCORECARD + [("合计", *[sum(r[i] for r in SCORECARD) for i in range(1, 6)])], [36, 12, 14, 12, 12, 12],
    "每项 1–5 分 × 权重；≥3.8 分进入投决讨论，任一一票否决项（接管率造假、关联收入 >50%、核心团队无量产经验且无补强计划）直接否决")
H_VC = ["环节", "作用", "成熟度", "供应商结构/集中度", "成本占比/价值量", "进入壁垒", "潜在利润率（判断）", "投资价值（判断）"]
add("09_产业链环节画像", H_VC, VALUE_CHAIN, [20, 20, 12, 34, 22, 24, 20, 20])

# 模型表
bom_rows = []
from models import BOM_FULL, ASSEMBLY  # noqa: E402
for name, q, lo, ba, hi, note in BOM_FULL:
    bom_rows.append((name, q, lo, ba, hi, lo * q, ba * q, hi * q, note))
add("10_BOM模型", ["部件", "数量", "单价低（元）", "单价基准（元）", "单价高（元）", "小计低", "小计基准", "小计高", "说明"], bom_rows,
    [44, 6, 12, 12, 12, 12, 12, 12, 40], f"装配测试良率加成：低 {ASSEMBLY['低']:.0%}/基准 {ASSEMBLY['基准']:.0%}/高 {ASSEMBLY['高']:.0%}；合计（元）：{MODEL['bom_full']}；三大件占物料 {MODEL['bom_big3_share']:.1%}（D 级假设，以公开锚点校验）")
tco_rows = [(r["方案"], round(r["初始投入"], 1), round(r["年化资本成本"], 2), round(r["年运营成本"], 2), round(r["替代人力当量"], 2),
             round(r["年节省人工"], 2), round(r["年净现金流(未含资本)"], 2), r["静态回本(年)"] and round(r["静态回本(年)"], 2), round(r["单位人力当量年成本"], 2))
            for r in MODEL["tco_factory"]]
add("11_工厂场景TCO", ["方案", "初始投入（万元）", "年化资本成本", "年运营成本", "替代人力当量", "年节省人工", "年净现金流", "静态回本（年）", "每FTE年化成本"], tco_rows,
    [36, 14, 14, 12, 12, 12, 12, 12, 14], "双班 4,800 小时/年；一线全口径用工成本 10 万元/人·年；资本成本 8%（D 级假设）")
sc_rows = []
for name, arr in MODEL["scenarios"].items():
    for e in arr:
        sc_rows.append((name, e["year"], e["units_wan"], e["asp_k_usd"], e["market_100m_usd"]))
add("12_三情景预测", ["情景", "年份", "全球出货（万台）", "均价（千美元）", "市场规模（亿美元）"], sc_rows, [10, 8, 16, 14, 18], "2025 为 IDC 实际值；其余为 D 级情景假设")
vc_rows = [(v["name"], v["post"], v["invest"], round(v["own0"] * 100, 2), round(v["own_exit"] * 100, 2), round(v["E_moic"], 2), round(v["irr"] * 100, 1)) for v in MODEL["vc"]]
add("13_VC回报模型", ["项目类型", "投后估值（亿元）", "投资（亿元）", "初始持股%", "退出持股%", "期望MOIC", "IRR%"], vc_rows, [40, 14, 12, 12, 12, 12, 10], "示例性假设，不代表任何真实企业")
mult_rows = [(m["name"], m["valuation_100m_cny"], m["revenue_100m_cny"], m["ev_s"] and round(m["ev_s"], 1)) for m in MODEL["multiples"]]
add("14_估值倍数横截面", ["公司/口径", "估值/市值（亿元）", "收入（亿元）", "估值/收入"], mult_rows, [32, 16, 14, 12], "汇率按 1 美元≈7.1 元人民币近似")

H_DD = ["编号", "类别", "问题", "为何问/验证的投资假设", "积极信号", "风险信号", "底层证据与独立验证"]
dd_rows = [(f"Q{i+1:02d}", *r) for i, r in enumerate(DD)]
add("15_尽调问题库", H_DD, dd_rows, [7, 10, 50, 34, 28, 28, 36], f"共 {len(DD)} 题；★ 为关键问题")
H_TH = ["编号", "结论", "核心事实", "机制", "反面证据/替代解释", "必要条件", "时间尺度", "领先指标", "影响", "投资指导", "推翻证据", "置信度"]
add("16_关键判断与反证", H_TH, THESES, [6, 40, 50, 36, 36, 28, 10, 30, 30, 34, 34, 8])
H_SRC = ["编号", "主题", "来源", "链接", "日期", "证据等级", "用途/口径说明"]
add("17_来源与证据索引", H_SRC, SOURCES, [6, 26, 22, 60, 12, 8, 40])

# ---------------- 写 xlsx ----------------
wb = Workbook()
wb.remove(wb.active)
head_fill = PatternFill("solid", fgColor="1F3864")
head_font = Font(bold=True, color="FFFFFF", name="Microsoft YaHei", size=10)
body_font = Font(name="Microsoft YaHei", size=9)
thin = Side(style="thin", color="D0D0D0")
readme = wb.create_sheet("00_说明")
readme_lines = [
    "具身智能全产业链深度研究与投资决策报告 · 数据底表",
    "信息截止：2026-10-21（部分事件更新至 2026-10-30）",
    "证据等级：A 一手资料或高质量独立验证；B 可靠二手来源且有交叉佐证；C 企业单方披露或单一来源；D 研究假设/推断/情景预测",
    "成熟度：L1 实验室论文与单场景演示；L2 受控条件持续运行；L3 真实客户现场试点；L4 可重复、可付费商业部署；L5 可规模化复制的成熟产品",
    "口径提醒：出货≠销量≠部署≠有效运行；意向订单≠正式订单≠确认收入≠回款；融资宣布≠交割；估值≠到账现金",
    "模型表（10–14）中的参数为示例性研究假设，不代表任何企业真实财务信息",
    "研究环境限制：多数网页无法直接打开原文，核验依赖检索摘录与多源比对；关键数字以交易所公告、企业正式披露为准",
]
for i, line in enumerate(readme_lines, 1):
    c = readme.cell(row=i, column=1, value=line)
    c.font = Font(name="Microsoft YaHei", size=11 if i > 1 else 14, bold=(i == 1))
readme.column_dimensions["A"].width = 120
for name, header, rows, widths, note in SHEETS:
    ws = wb.create_sheet(name[:31])
    start = 1
    if note:
        ws.cell(row=1, column=1, value=note).font = Font(name="Microsoft YaHei", size=9, italic=True, color="555555")
        start = 2
    for j, h in enumerate(header, 1):
        c = ws.cell(row=start, column=j, value=h)
        c.fill, c.font = head_fill, head_font
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for i, r in enumerate(rows, start + 1):
        for j, v in enumerate(r, 1):
            c = ws.cell(row=i, column=j, value=v)
            c.font = body_font
            c.alignment = Alignment(wrap_text=True, vertical="top")
            c.border = Border(bottom=thin)
    if widths:
        for j, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = ws.cell(row=start + 1, column=2)
    ws.auto_filter.ref = f"A{start}:{get_column_letter(len(header))}{start + len(rows)}"
    # csv
    with open(os.path.join(DATA, f"{name}.csv"), "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(header)
        w.writerows(rows)
xlsx_path = os.path.join(DATA, "具身智能投资研究_数据底表_2026-10.xlsx")
wb.save(xlsx_path)

# ---------------- 附录 Markdown 片段 ----------------
save_md("a1_companies", md_table(["名称", "地区", "环节", "资本状态", "最新估值/市值", "收入/出货", "成熟度", "证据"],
                                 [(c[0], c[1], c[2], c[4], c[5], c[6], c[7], c[8]) for c in COMPANIES]))
save_md("a1_companies_risk", md_table(["名称", "核心产品/路线", "关键风险/备注"], [(c[0], c[3], c[9]) for c in COMPANIES]))
save_md("a2_fin_events", md_table(H_FIN, FIN_EVENTS))
save_md("a2_fin_stats", md_table(H_STAT, FIN_STATS))
save_md("a2_shipments", md_table(H_SHIP, SHIPMENTS))
save_md("a3_tech", md_table(H_TECH, TECH_MATRIX))
save_md("a4_scen", md_table(H_SCEN, SCENARIOS))
save_md("a6_opp", md_table(H_OPP, OPPORTUNITIES))
save_md("a_scorecard", md_table(H_SC, list(SCORECARD) + [("合计", *[sum(r[i] for r in SCORECARD) for i in range(1, 6)])]))
save_md("a_valuechain", md_table(H_VC, VALUE_CHAIN))
save_md("a8_dd", md_table(["编号", "类别", "问题", "验证的假设", "积极信号", "风险信号", "证据与验证"], dd_rows))
save_md("a9_sources", md_table(["编号", "主题", "来源", "链接", "日期", "等级", "用途/口径"], SOURCES))
save_md("a7_theses_short", md_table(["编号", "结论", "领先指标", "推翻证据", "置信度"], [(t[0], t[1], t[7], t[10], t[11]) for t in THESES]))

# 第十六篇正文：命题完整展开
blocks = []
for t in THESES:
    (no, concl, facts, mech, counter, cond, horizon, ind, impact, guide, falsify, conf) = t
    blocks.append(f"### {no}　{concl}\n")
    blocks.append(f"- **核心事实**：{facts}")
    blocks.append(f"- **机制**：{mech}")
    blocks.append(f"- **反面证据/替代解释**：{counter}")
    blocks.append(f"- **成立的必要条件**：{cond}")
    blocks.append(f"- **时间尺度**：{horizon}　｜　**置信度**：{conf}")
    blocks.append(f"- **领先指标**：{ind}")
    blocks.append(f"- **对环节与企业的影响**：{impact}")
    blocks.append(f"- **投资指导**：{guide}")
    blocks.append(f"- **推翻证据**：{falsify}\n")
save_md("c16_theses", "\n".join(blocks))

# 关键尽调问题（★）展开
key = [r for r in dd_rows if r[2].startswith("★")]
kb = []
for r in key:
    kb.append(f"**{r[0]}（{r[1]}）{r[2].lstrip('★ ').strip()}**\n")
    kb.append(f"- 为何必须问／验证的假设：{r[3]}")
    kb.append(f"- 积极信号：{r[4]}")
    kb.append(f"- 风险信号：{r[5]}")
    kb.append(f"- 底层证据与独立验证：{r[6]}\n")
save_md("c15_key_dd", "\n".join(kb))

print("sheets:", len(SHEETS), "companies:", len(COMPANIES), "events:", len(FIN_EVENTS), "dd:", len(DD), "key dd:", len(key), "theses:", len(THESES), "sources:", len(SOURCES))
print(xlsx_path)
